"""Ophalen en parsen van de gratis, publieke bronnen. Geen API-sleutels nodig.

Elke functie geeft een reeks terug: [(\"YYYY-MM-DD\", float), ...] oplopend op datum.
De parse_*-functies zijn los testbaar zonder netwerk.
"""
import csv
import io
import re
import time
from datetime import date, datetime, timezone

import requests

UA = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 Chrome/126 Safari/537.36"}
START = "1990-01-01"
TIMEOUT = 30


def _get(url, params=None, tries=3):
    last = None
    for i in range(tries):
        try:
            r = requests.get(url, params=params, headers=UA, timeout=TIMEOUT)
            r.raise_for_status()
            return r
        except requests.RequestException as e:  # netwerkfout of 4xx/5xx
            last = e
            time.sleep(2 * (i + 1))
    raise RuntimeError(f"Ophalen mislukt: {url} ({last})")


# ---------- FRED (CSV-download, zonder sleutel) ----------
def parse_fred_csv(text):
    out = []
    rows = csv.reader(io.StringIO(text))
    next(rows, None)
    for row in rows:
        if len(row) < 2 or row[1] in ("", "."):
            continue
        try:
            out.append((row[0], float(row[1])))
        except ValueError:
            continue
    return out


def fred(series_id, start=START):
    r = _get("https://fred.stlouisfed.org/graph/fredgraph.csv", {"id": series_id, "cosd": start})
    data = parse_fred_csv(r.text)
    if not data:
        raise RuntimeError(f"FRED gaf geen data voor {series_id}")
    return data


# ---------- Yahoo Finance (chart-endpoint) ----------
def parse_yahoo_chart(payload):
    res = payload["chart"]["result"][0]
    closes = res["indicators"]["quote"][0]["close"]
    out = {}
    for ts, c in zip(res.get("timestamp", []), closes):
        if c is None:
            continue
        d = datetime.fromtimestamp(ts, tz=timezone.utc).date().isoformat()
        out[d] = float(c)
    return sorted(out.items())


def yahoo(symbol, start=START):
    p1 = int(datetime.fromisoformat(start).replace(tzinfo=timezone.utc).timestamp())
    r = _get(
        f"https://query1.finance.yahoo.com/v8/finance/chart/{requests.utils.quote(symbol)}",
        {"period1": p1, "period2": int(time.time()), "interval": "1d"},
    )
    data = parse_yahoo_chart(r.json())
    if not data:
        raise RuntimeError(f"Yahoo gaf geen data voor {symbol}")
    return data


# ---------- multpl.com (HTML-tabel) ----------
_ROW = re.compile(
    r"<td[^>]*>\s*([A-Z][a-z]{2} \d{1,2}, \d{4})\s*</td>\s*<td[^>]*>(.*?)</td>", re.S
)
_NUM = re.compile(r"-?\d[\d,]*\.?\d*")


def parse_multpl(html):
    out = {}
    for d, cell in _ROW.findall(html):
        m = _NUM.search(re.sub(r"<[^>]+>|&[#\w]+;", " ", cell))
        if not m:
            continue
        dt = datetime.strptime(d, "%b %d, %Y").date()
        # één waarde per maand; de tabel loopt van nieuw naar oud, dus de eerste wint
        out.setdefault(dt.strftime("%Y-%m-01"), float(m.group().replace(",", "")))
    return sorted(out.items())


def multpl(path, start=START):
    r = _get(f"https://www.multpl.com/{path}/table/by-month")
    data = [x for x in parse_multpl(r.text) if x[0] >= start]
    if not data:
        raise RuntimeError(f"multpl gaf geen data voor {path}")
    return data


# ---------- FINRA margin statistics (xlsx) ----------
FINRA_XLSX = "https://www.finra.org/sites/default/files/2021-03/margin-statistics.xlsx"


def _finra_date(v):
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-01")
    s = str(v).strip()
    for fmt in ("%b-%y", "%b-%Y", "%Y-%m", "%B %Y", "%b %Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-01")
        except ValueError:
            pass
    return None


def parse_finra_xlsx(content):
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    col, out = None, {}
    for row in ws.iter_rows(values_only=True):
        if col is None:
            for i, c in enumerate(row):
                if c and "debit" in str(c).lower():
                    col = i
            continue
        if not row or row[0] is None or col >= len(row) or row[col] is None:
            continue
        d = _finra_date(row[0])
        try:
            val = float(str(row[col]).replace(",", ""))
        except ValueError:
            continue
        if d:
            out[d] = val
    return sorted(out.items())


def finra_margin_debt():
    data = parse_finra_xlsx(_get(FINRA_XLSX).content)
    if not data:
        raise RuntimeError("FINRA-bestand bevat geen herkenbare data")
    return data


if __name__ == "__main__":  # snelle handmatige controle: python -m app.sources
    for name, fn in [("FRED T10Y2Y", lambda: fred("T10Y2Y")), ("Yahoo ^VIX", lambda: yahoo("^VIX")),
                     ("multpl CAPE", lambda: multpl("shiller-pe")), ("FINRA", finra_margin_debt)]:
        try:
            s = fn()
            print(f"OK   {name}: {len(s)} punten, laatste {s[-1]}")
        except Exception as e:
            print(f"FOUT {name}: {e}")
